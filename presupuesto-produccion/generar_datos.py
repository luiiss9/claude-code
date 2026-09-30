#!/usr/bin/env python3
"""Genera la base de datos de producción 2025 (Excel + CSV) y la incrusta como
datos de ejemplo en presupuesto_produccion_2026.html.

Uso:  python3 generar_datos.py
Requiere: pip install openpyxl
"""
import csv
import json
import os
import random
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(2025)
HERE = os.path.dirname(os.path.abspath(__file__))
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
SEASON = [0.93, 0.96, 1.03, 1.00, 1.05, 1.04, 0.92, 0.62, 1.04, 1.10, 1.12, 1.07]

# (columna, descripción, grupo de la diapositiva, tipo F/V)
CONCEPTOS = [
    ("Materia_Prima", "Materia prima principal", "Consumos directos de materiales", "V"),
    ("Materiales_Auxiliares", "Adhesivos, embalajes, repuestos específicos", "Consumos directos de materiales", "V"),
    ("Componentes_Semielaborados", "Componentes o semielaborados adquiridos a terceros", "Consumos directos de materiales", "V"),
    ("Material_Indirecto", "Lubricantes, herramientas pequeñas, suministros de taller", "Consumos directos de materiales", "V"),
    ("MOD", "Mano de obra directa (operarios de línea)", "Costes de mano de obra", "V"),
    ("MOI", "Mano de obra indirecta (supervisores, mantenimiento básico)", "Costes de mano de obra", "F"),
    ("Cargas_Sociales", "Cargas sociales y cotizaciones asociadas", "Costes de mano de obra", "V"),
    ("Energia_Suministros", "Energía eléctrica, gas, vapor, agua, aire comprimido", "Costes indirectos de fabricación (CIF)", "V"),
    ("Mantenimiento", "Mantenimiento preventivo y correctivo de maquinaria", "Costes indirectos de fabricación (CIF)", "F"),
    ("Amortizacion", "Amortización de maquinaria, instalaciones y utillaje", "Costes indirectos de fabricación (CIF)", "F"),
    ("Arrendamientos", "Arrendamientos o cánones de equipos de producción", "Costes indirectos de fabricación (CIF)", "F"),
    ("Servicios_Externos", "Servicios externos ligados a producción", "Costes indirectos de fabricación (CIF)", "F"),
    ("Seguros", "Seguros de maquinaria e instalaciones productivas", "Costes indirectos de fabricación (CIF)", "F"),
    ("Calidad_Control", "Costes de calidad y control de procesos", "Costes indirectos de fabricación (CIF)", "V"),
    ("Transporte_Interno", "Carretillas, AGVs, cintas", "Costes logísticos ligados a producción", "V"),
    ("Manipulacion_Almacenaje", "Manipulación y almacenaje de materias primas y semielaborados", "Costes logísticos ligados a producción", "V"),
    ("Embalajes_Industriales", "Palets, cajas, protecciones", "Costes logísticos ligados a producción", "V"),
    ("Distribucion_Interna", "Distribución interna hacia otras secciones o centros", "Costes logísticos ligados a producción", "V"),
    ("Productos_Defectuosos", "Coste de productos defectuosos (rechazos internos)", "Desperdicios y mermas", "V"),
    ("Mermas_Materiales", "Pérdidas por mermas de materiales en el proceso", "Desperdicios y mermas", "V"),
    ("Residuos_Industriales", "Tratamiento y gestión de residuos industriales", "Desperdicios y mermas", "V"),
    ("Seguridad_Higiene", "Seguridad e higiene en planta (EPI, revisiones)", "Otros costes asociados a producción", "F"),
    ("Formacion", "Formación específica del personal productivo", "Otros costes asociados a producción", "F"),
    ("Innovacion_Mejora", "Innovación y mejora continua aplicada a fabricación", "Otros costes asociados a producción", "F"),
]

# nombre, familia, planta, uds/mes base, coste materia prima €/ud, horas MOD/ud, horas máq/ud, tasa defectos, factor capacidad
PRODUCTOS = [
    ("Bomba Hidraulica BH-200", "Hidraulica", "Planta Norte", 1800, 46.0, 0.55, 0.40, 0.030, 1.28),
    ("Valvula Reguladora VR-50", "Hidraulica", "Planta Norte", 3200, 14.5, 0.22, 0.18, 0.024, 1.30),
    ("Motor Electrico ME-15", "Electrica", "Planta Norte", 1100, 88.0, 0.80, 0.55, 0.035, 1.20),
    ("Cuadro Electrico CE-8", "Electrica", "Planta Norte", 950, 61.0, 0.95, 0.30, 0.028, 1.24),
    ("Reductor RD-30", "Mecanica", "Planta Sur", 1500, 52.0, 0.60, 0.48, 0.032, 1.26),
    ("Compresor CP-10", "Mecanica", "Planta Sur", 700, 132.0, 1.30, 0.85, 0.040, 1.18),
    ("Kit Repuestos KR-01", "Repuestos", "Planta Sur", 5200, 6.8, 0.09, 0.07, 0.018, 1.32),
    ("Filtro Industrial FI-5", "Repuestos", "Planta Sur", 4100, 9.2, 0.12, 0.10, 0.022, 1.30),
]
TARIFA_MOD = 17.8      # €/hora incl. salario bruto
TARIFA_ENERGIA = 5.6   # €/hora máquina
mean_season = sum(SEASON) / 12


def ruido(p=0.03):
    return 1 + random.uniform(-p, p)


COLS = ["Año", "Mes", "Numero_Mes", "Planta", "Familia", "Producto",
        "Unidades_Producidas", "Unidades_Defectuosas", "Horas_Maquina", "Capacidad_Max_Unidades"] \
    + [c[0] for c in CONCEPTOS]

rows = []
for (nom, fam, planta, base, mp, hmod, hmaq, tdef, capf) in PRODUCTOS:
    fijo_base = base * mp
    fijos = {
        "MOI": 0.055 * fijo_base, "Mantenimiento": 0.030 * fijo_base,
        "Amortizacion": 0.075 * fijo_base, "Arrendamientos": 0.022 * fijo_base,
        "Servicios_Externos": 0.020 * fijo_base, "Seguros": 0.008 * fijo_base,
        "Seguridad_Higiene": 0.008 * fijo_base, "Formacion": 0.006 * fijo_base,
        "Innovacion_Mejora": 0.010 * fijo_base,
    }
    for m in range(12):
        uds = round(base * SEASON[m] / mean_season * (1 + 0.004 * m) * ruido(0.03))
        defu = round(uds * tdef * ruido(0.15))
        horas_maq = round(uds * hmaq * ruido(0.02), 1)
        mod = uds * hmod * TARIFA_MOD * ruido(0.02)
        r = {
            "Año": 2025, "Mes": MESES[m], "Numero_Mes": m + 1, "Planta": planta,
            "Familia": fam, "Producto": nom, "Unidades_Producidas": uds,
            "Unidades_Defectuosas": defu, "Horas_Maquina": horas_maq,
            "Capacidad_Max_Unidades": round(base * capf),
            "Materia_Prima": uds * mp * ruido(0.02),
            "Materiales_Auxiliares": uds * mp * 0.06 * ruido(0.03),
            "Componentes_Semielaborados": uds * mp * 0.25 * ruido(0.03),
            "Material_Indirecto": uds * mp * 0.02 * ruido(0.04),
            "MOD": mod,
            "Cargas_Sociales": mod * 0.31,
            "Energia_Suministros": horas_maq * TARIFA_ENERGIA * ruido(0.04),
            "Calidad_Control": uds * mp * 0.015 * ruido(0.04),
            "Transporte_Interno": uds * mp * 0.012 * ruido(0.05),
            "Manipulacion_Almacenaje": uds * mp * 0.020 * ruido(0.04),
            "Embalajes_Industriales": uds * mp * 0.030 * ruido(0.03),
            "Distribucion_Interna": uds * mp * 0.015 * ruido(0.05),
            "Productos_Defectuosos": defu * (mp * 1.35 + hmod * TARIFA_MOD) * 0.55,
            "Mermas_Materiales": uds * mp * 0.015 * ruido(0.08),
            "Residuos_Industriales": uds * mp * 0.004 * ruido(0.08),
        }
        for k, v in fijos.items():
            r[k] = v * ruido(0.015)
        for c in COLS:
            if isinstance(r[c], float):
                r[c] = round(r[c], 2)
        rows.append(r)

assert len(rows) <= 500

# ---------- CSV ----------
csv_path = os.path.join(HERE, "datos_produccion_2025.csv")
with open(csv_path, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=COLS)
    w.writeheader()
    w.writerows(rows)

# ---------- Excel ----------
wb = Workbook()
ws = wb.active
ws.title = "Produccion_2025"
hdr_fill = PatternFill("solid", fgColor="1F3864")
grp_colors = {
    "Consumos directos de materiales": "2E75B6", "Costes de mano de obra": "548235",
    "Costes indirectos de fabricación (CIF)": "BF8F00", "Costes logísticos ligados a producción": "C55A11",
    "Desperdicios y mermas": "C00000", "Otros costes asociados a producción": "7030A0",
}
grp_of = {c[0]: c[2] for c in CONCEPTOS}
ws.append(COLS)
for r in rows:
    ws.append([r[c] for c in COLS])
for i, c in enumerate(COLS, 1):
    cell = ws.cell(row=1, column=i)
    color = grp_colors.get(grp_of.get(c), "1F3864")
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = max(12, min(28, len(c) + 4))
    if i >= 7:
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=i).number_format = "#,##0" if i <= 10 else "#,##0.00"
ws.row_dimensions[1].height = 36
ws.freeze_panes = "G2"
ws.auto_filter.ref = ws.dimensions

wd = wb.create_sheet("Diccionario")
wd.append(["Columna", "Descripción", "Grupo (presupuesto de producción)", "Tipo sugerido (F=fijo, V=variable)"])
base_desc = [
    ("Año", "Ejercicio", "Identificación", ""), ("Mes", "Nombre del mes", "Identificación", ""),
    ("Numero_Mes", "Mes 1-12 (opcional)", "Identificación", ""), ("Planta", "Centro de producción", "Identificación", ""),
    ("Familia", "Familia de productos", "Identificación", ""), ("Producto", "Producto o referencia", "Identificación", ""),
    ("Unidades_Producidas", "Unidades fabricadas en el mes", "Volumen", ""),
    ("Unidades_Defectuosas", "Unidades rechazadas (opcional)", "Volumen", ""),
    ("Horas_Maquina", "Horas máquina consumidas (opcional)", "Volumen", ""),
    ("Capacidad_Max_Unidades", "Capacidad máxima mensual del producto (opcional)", "Volumen", ""),
]
for r in base_desc + CONCEPTOS:
    wd.append(list(r))
for c in wd[1]:
    c.fill = hdr_fill
    c.font = Font(bold=True, color="FFFFFF")
for col, wdt in zip("ABCD", (28, 62, 42, 30)):
    wd.column_dimensions[col].width = wdt

wi = wb.create_sheet("Instrucciones")
for line in [
    "BASE DE DATOS DE PRODUCCIÓN 2025 (datos ficticios de ejemplo)",
    "",
    f"- Hoja 'Produccion_2025': {len(rows)} filas = {len(PRODUCTOS)} productos x 12 meses. Una fila por producto y mes.",
    "- Las 24 columnas de coste corresponden a las partidas de las diapositivas del Tema 6 (Control de Gestión):",
    "  materiales, mano de obra, CIF, logística, desperdicios/mermas y otros costes.",
    "- Sube este archivo (o el CSV) a presupuesto_produccion_2026.html para generar el presupuesto de 2026.",
    "- Puedes sustituir los datos por los reales manteniendo los nombres de columna; el HTML permite remapear columnas.",
    "- No añadas filas de totales dentro de la hoja de datos.",
]:
    wi.append([line])
wi["A1"].font = Font(bold=True, size=14)
wi.column_dimensions["A"].width = 120

xlsx_path = os.path.join(HERE, "datos_produccion_2025.xlsx")
wb.save(xlsx_path)

# ---------- Incrustar CSV de ejemplo en el HTML ----------
html_path = os.path.join(HERE, "presupuesto_produccion_2026.html")
if os.path.exists(html_path):
    with open(csv_path, encoding="utf-8") as f:
        sample = f.read()
    with open(html_path, encoding="utf-8") as f:
        html = f.read()
    new = "/*SAMPLE_START*/const SAMPLE_CSV = " + json.dumps(sample) + ";/*SAMPLE_END*/"
    html = re.sub(r"/\*SAMPLE_START\*/.*?/\*SAMPLE_END\*/", lambda _: new, html, flags=re.S)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)

print(f"OK: {len(rows)} filas, {len(COLS)} columnas -> {xlsx_path}")
